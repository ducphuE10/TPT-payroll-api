from io import BytesIO
import os
from docx import Document
from fastapi import Depends
from fastapi.security.utils import get_authorization_scheme_param
from jose import JWTError, jwt
from jose.exceptions import JWKError
from payroll.config import settings
import logging
from typing import Annotated
from fastapi.security import APIKeyHeader
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from payroll.dependants.repositories import (
    # retrieve_dependant_by_cccd,
    retrieve_dependant_by_mst,
)
from payroll.employees.repositories import (
    retrieve_employee_by_cccd,
    retrieve_employee_by_mst,
)

logger = logging.getLogger(__name__)


api_key_header = APIKeyHeader(name="Authorization")
TokenDep = Annotated[str, Depends(api_key_header)]


def get_user_email(authorization, **kwargs):
    scheme, param = get_authorization_scheme_param(authorization)
    if not authorization or scheme.lower() != "bearer":
        logger.exception(
            f"Malformed authorization header. Scheme: {scheme} Param: {param} Authorization: {authorization}"
        )
        return

    token = authorization.split()[1]

    try:
        data = jwt.decode(token, settings.SECRET_KEY.get_secret_value())
    except (JWKError, JWTError):
        raise Exception("Invalid token")
    return data["email"]


def get_error_message_dict():
    return {
        "SYSTEM_EXCEPTION": {"ERR_SM_99999": "System error, please try again later."},
        "APP_EXCEPTION": {
            "ERR_INVALID_INPUT": "Invalid input data.",
            "ERR_RESOURCE_NOT_FOUND": "Resource not found.",
            "ERR_RESOURCE_CONFLICT": "Resource conflict.",
            "ERR_RESOURCE_ALREADY_EXISTS": "Resource already exists.",
            "ERR_FORBIDDEN_ACTION": "Forbidden action.",
            "ERR_USER_WITH_EMAIL_ALREADY_EXISTS": "User with email already exists.",
            "ERR_INVALID_USERNAME_OR_PASSWORD": "Invalid username or password.",
            "ERR_CANNOT_CREATE_ADMIN_USER": "Cannot create admin user.",
            "ERR_EXIST_DEPEND_OBJECT": "Cannot delete due to related object existing.",
            "ERR_WORK_LEAVE_STATE": "Invalid work or leave state.",
        },
    }


def check_exist_person_by_cccd(
    *,
    db_session,
    cccd: str,
    exclude_id: int = None,
):
    employee = retrieve_employee_by_cccd(
        db_session=db_session,
        employee_cccd=cccd,
        exclude_employee_id=exclude_id,
    )

    return bool(employee)


def check_exist_person_by_mst(
    *,
    db_session,
    mst: str,
    exclude_id: int = None,
):
    employee = retrieve_employee_by_mst(
        db_session=db_session,
        employee_mst=mst,
        exclude_employee_id=exclude_id,
    )
    dependant = retrieve_dependant_by_mst(
        db_session=db_session,
        dependant_mst=mst,
        exclude_dependant_id=exclude_id,
    )
    return bool(employee or dependant)


def fill_template(template_path: str, data: dict):
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found at {template_path}")
    doc = Document(template_path)
    # Iterate over paragraphs and replace placeholders with actual data

    def set_font_to_times_new_roman(paragraph):
        """Sets the font of a paragraph to Times New Roman."""
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(11)

    def replace_text_in_paragraph(paragraph, data):
        """Replaces placeholders in the entire paragraph text."""
        # Combine the text from all runs into one string
        full_text = "".join([run.text for run in paragraph.runs])
        # Replace placeholders with actual values
        for key, value in data.items():
            placeholder = f"{{{{ {key} }}}}"
            if placeholder in full_text:
                full_text = full_text.replace(placeholder, value)

        # Set the replaced text back into the runs
        if full_text:
            paragraph.runs[0].text = full_text  # Assign all text to the first run
            for run in paragraph.runs[1:]:  # Clear text in remaining runs
                run.text = ""

        set_font_to_times_new_roman(paragraph)

    def replace_text_in_table(table, data):
        """Replaces placeholders in all cells of a table."""
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_text_in_paragraph(paragraph, data)

    # Iterate over all elements in the document in the correct order
    def iter_block_items(parent):
        """Yield paragraphs and tables in the order they appear in the document."""
        for child in parent.element.body:
            if child.tag.endswith("p"):
                yield Paragraph(child, parent)
            elif child.tag.endswith("tbl"):
                yield Table(child, parent)

    # Process paragraphs and tables in the document in the correct order
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            replace_text_in_paragraph(block, data)
        elif isinstance(block, Table):
            replace_text_in_table(block, data)

    # Save the modified document to a bytes buffer
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    return file_stream


def format_with_dot(value):
    return f"{value:,.0f}".replace(",", ".")
